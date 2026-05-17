#!/usr/bin/env python3
"""
BAND-AID 6 reference extractor.

Reads references/references.json, extracts text from each source file
(PDF via PyMuPDF, PPTX via python-pptx), chunks it, and writes
references/extracted/<id>.json sidecars used by the in-browser RAG retrieval.

Run from repo root:
    python3 tools/extract.py

Re-run any time you add or change a source file.

Output schema (per file):
{
  "id": "<reference id>",
  "title": "<reference title>",
  "category": "<reference category>",
  "type": "<pdf|presentation|image|...>",
  "chunks": [
    { "id": "<refId>#<n>", "section": "p.3" | "Slide 2", "text": "..." },
    ...
  ]
}

Image references (and any type we cannot text-extract) are skipped.
"""

import json
import os
import re
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
MANIFEST = REPO / "references" / "references.json"
OUT_DIR = REPO / "references" / "extracted"
SOURCE_DIR = REPO / "references" / "source"

# Chunking: target ~700 chars, hard cap 1200, soft min 200.
TARGET = 700
HARD_MAX = 1200
SOFT_MIN = 200

WHITESPACE = re.compile(r"[ \t]+")
MULTI_NEWLINE = re.compile(r"\n{3,}")


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # collapse runs of spaces/tabs but keep newlines for paragraph hints
    text = "\n".join(WHITESPACE.sub(" ", line).strip() for line in text.split("\n"))
    text = MULTI_NEWLINE.sub("\n\n", text).strip()
    return text


def chunk_paragraphs(text: str, section: str, ref_id: str, start_idx: int):
    """Greedy paragraph-aware chunker. Returns (chunks, next_idx)."""
    chunks = []
    paras = [p.strip() for p in text.split("\n\n") if p.strip()]
    buf = ""
    idx = start_idx
    for p in paras:
        if not buf:
            buf = p
            continue
        # If adding this paragraph stays under HARD_MAX, append it.
        if len(buf) + 2 + len(p) <= HARD_MAX:
            buf = buf + "\n\n" + p
            # If we have already passed TARGET, flush.
            if len(buf) >= TARGET:
                chunks.append({"id": f"{ref_id}#{idx}", "section": section, "text": buf})
                idx += 1
                buf = ""
        else:
            # Flush current buf, start a new one with p.
            chunks.append({"id": f"{ref_id}#{idx}", "section": section, "text": buf})
            idx += 1
            # If p itself is larger than HARD_MAX, hard-split it.
            if len(p) > HARD_MAX:
                for i in range(0, len(p), HARD_MAX):
                    piece = p[i : i + HARD_MAX]
                    chunks.append({"id": f"{ref_id}#{idx}", "section": section, "text": piece})
                    idx += 1
                buf = ""
            else:
                buf = p
    if buf:
        # Merge tiny tail into prior chunk if possible.
        if chunks and len(buf) < SOFT_MIN and len(chunks[-1]["text"]) + 2 + len(buf) <= HARD_MAX and chunks[-1]["section"] == section:
            chunks[-1]["text"] = chunks[-1]["text"] + "\n\n" + buf
        else:
            chunks.append({"id": f"{ref_id}#{idx}", "section": section, "text": buf})
            idx += 1
    return chunks, idx


def extract_pdf(path: Path, ref_id: str):
    import fitz
    doc = fitz.open(str(path))
    chunks = []
    idx = 1
    for page_num, page in enumerate(doc, start=1):
        raw = page.get_text("text")
        clean = normalize(raw)
        if not clean:
            continue
        page_chunks, idx = chunk_paragraphs(clean, f"p.{page_num}", ref_id, idx)
        chunks.extend(page_chunks)
    doc.close()
    return chunks


def extract_pptx(path: Path, ref_id: str):
    from pptx import Presentation
    p = Presentation(str(path))
    chunks = []
    idx = 1
    for slide_num, slide in enumerate(p.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        parts.append(t)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip():
                        parts.append(row_text)
        raw = "\n\n".join(parts)
        clean = normalize(raw)
        if not clean:
            continue
        slide_chunks, idx = chunk_paragraphs(clean, f"Slide {slide_num}", ref_id, idx)
        chunks.extend(slide_chunks)
    return chunks


def main() -> int:
    if not MANIFEST.exists():
        print(f"ERROR: manifest not found at {MANIFEST}", file=sys.stderr)
        return 1
    manifest = json.loads(MANIFEST.read_text())
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    index = []
    total_chunks = 0
    for entry in manifest:
        ref_id = entry["id"]
        ref_type = entry.get("type", "")
        src_path = REPO / entry["path"]
        title = entry.get("title", ref_id)
        category = entry.get("category", "")
        print(f"→ {ref_id} ({ref_type}) {src_path.name}")
        if not src_path.exists():
            print(f"   SKIP — source file missing: {src_path}")
            continue
        try:
            if ref_type == "pdf":
                chunks = extract_pdf(src_path, ref_id)
            elif ref_type == "presentation":
                chunks = extract_pptx(src_path, ref_id)
            else:
                print(f"   SKIP — type '{ref_type}' has no text extractor")
                continue
        except Exception as e:
            print(f"   ERROR extracting: {e}")
            continue

        out = {
            "id": ref_id,
            "title": title,
            "category": category,
            "type": ref_type,
            "chunks": chunks,
        }
        out_path = OUT_DIR / f"{ref_id}.json"
        out_path.write_text(json.dumps(out, ensure_ascii=False, indent=2))
        total_chars = sum(len(c["text"]) for c in chunks)
        print(f"   {len(chunks)} chunks, {total_chars} chars → {out_path.relative_to(REPO)}")
        index.append({"id": ref_id, "title": title, "category": category, "type": ref_type,
                      "chunk_count": len(chunks), "char_count": total_chars,
                      "file": f"references/extracted/{ref_id}.json"})
        total_chunks += len(chunks)

    index_path = OUT_DIR / "index.json"
    index_path.write_text(json.dumps(index, ensure_ascii=False, indent=2))
    print(f"\nWrote index with {len(index)} entries and {total_chunks} total chunks → {index_path.relative_to(REPO)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
